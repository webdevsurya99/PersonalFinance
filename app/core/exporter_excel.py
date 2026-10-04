import io
import datetime
from typing import List
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.models import Transaction, User

def generate_excel_statement(transactions: List[Transaction], user: User, start_date: datetime.date, end_date: datetime.date) -> io.BytesIO:
    """
    Generate a beautifully formatted Excel statement for the user's transactions.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Statement of Accounts"

    # Color definitions
    PRIMARY_COLOR = "1E293B"      # Dark slate
    HEADER_TEXT_COLOR = "FFFFFF"
    EXPENSE_COLOR = "EF4444"      # Red
    INCOME_COLOR = "10B981"       # Green
    LIGHT_BG = "F8FAFC"
    ALT_ROW_BG = "F1F5F9"

    # Styles
    title_font = Font(name="Segoe UI", size=16, bold=True, color="0F172A")
    subtitle_font = Font(name="Segoe UI", size=10, italic=True, color="64748B")
    header_font = Font(name="Segoe UI", size=11, bold=True, color=HEADER_TEXT_COLOR)
    header_fill = PatternFill(start_color=PRIMARY_COLOR, end_color=PRIMARY_COLOR, fill_type="solid")
    
    bold_font = Font(name="Segoe UI", size=10, bold=True)
    regular_font = Font(name="Segoe UI", size=10)
    income_font = Font(name="Segoe UI", size=10, bold=True, color="047857")
    expense_font = Font(name="Segoe UI", size=10, bold=True, color="B91C1C")

    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    # 1. Header Information
    ws.merge_cells("A1:G1")
    ws["A1"] = f"Personal Financial Statement - {user.full_name}"
    ws["A1"].font = title_font
    ws["A1"].alignment = Alignment(vertical="center")

    ws.merge_cells("A2:G2")
    ws["A2"] = f"Period: {start_date.strftime('%d %B %Y')} to {end_date.strftime('%d %B %Y')}  |  Generated on: {datetime.datetime.now().strftime('%d-%m-%Y %H:%M')}"
    ws["A2"].font = subtitle_font
    ws["A2"].alignment = Alignment(vertical="center")

    # 2. KPI Summary Box
    total_income = sum(t.amount for t in transactions if t.type == "income")
    total_expense = sum(t.amount for t in transactions if t.type == "expense")
    net_savings = total_income - total_expense

    ws["A4"] = "Total Income"
    ws["B4"] = f"{user.currency_symbol} {total_income:,.2f}"
    ws["A4"].font = bold_font
    ws["B4"].font = income_font

    ws["C4"] = "Total Expenses"
    ws["D4"] = f"{user.currency_symbol} {total_expense:,.2f}"
    ws["C4"].font = bold_font
    ws["D4"].font = expense_font

    ws["E4"] = "Net Savings"
    ws["F4"] = f"{user.currency_symbol} {net_savings:,.2f}"
    ws["E4"].font = bold_font
    ws["F4"].font = income_font if net_savings >= 0 else expense_font

    for col in range(1, 7):
        cell = ws.cell(row=4, column=col)
        cell.fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
        cell.border = thin_border

    # 3. Transaction Table Headers
    headers = [
        "Date", "Type", "Category", "Payment Source / Destination", 
        "Trip Tag", "Tag", "Amount", "Remarks"
    ]
    
    start_row = 6
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=start_row, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
    
    ws.row_dimensions[start_row].height = 24

    # 4. Populate Transaction Rows
    current_row = start_row + 1
    for t in sorted(transactions, key=lambda x: x.date, reverse=True):
        source = "—"
        if t.bank_account:
            source = f"🏦 {t.bank_account.bank_name} (••{t.bank_account.account_number_last4})"
        elif t.credit_card:
            source = f"💳 {t.credit_card.card_name} (••{t.credit_card.last_4_digits})"

        cat_str = f"{t.category.icon} {t.category.name}" if t.category else "Uncategorized"
        trip_str = f"✈️ {t.trip.name}" if t.trip else "—"

        ws.cell(row=current_row, column=1, value=t.date.strftime("%d-%m-%Y")).alignment = Alignment(horizontal="center")
        
        type_cell = ws.cell(row=current_row, column=2, value=t.type.capitalize())
        type_cell.alignment = Alignment(horizontal="center")
        type_cell.font = income_font if t.type == "income" else expense_font

        ws.cell(row=current_row, column=3, value=cat_str)
        ws.cell(row=current_row, column=4, value=source)
        ws.cell(row=current_row, column=5, value=trip_str)
        ws.cell(row=current_row, column=6, value=t.tag or "—")

        amt_cell = ws.cell(
            row=current_row, 
            column=7, 
            value=f"{'+' if t.type == 'income' else '-'} {user.currency_symbol} {t.amount:,.2f}"
        )
        amt_cell.alignment = Alignment(horizontal="right")
        amt_cell.font = income_font if t.type == "income" else expense_font

        ws.cell(row=current_row, column=8, value=t.remarks or "")

        # Format and border
        is_alt = (current_row % 2 == 0)
        row_fill = PatternFill(start_color=ALT_ROW_BG, end_color=ALT_ROW_BG, fill_type="solid") if is_alt else PatternFill(fill_type=None)
        
        for c in range(1, 9):
            c_cell = ws.cell(row=current_row, column=c)
            if not c_cell.font.bold:
                c_cell.font = regular_font
            if is_alt and not c_cell.fill.fill_type:
                c_cell.fill = row_fill
            c_cell.border = thin_border

        ws.row_dimensions[current_row].height = 20
        current_row += 1

    # Auto-fit column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or "")
            max_len = max(max_len, len(val))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # Save to BytesIO
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
